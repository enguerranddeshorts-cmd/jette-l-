#!/usr/bin/env python3
"""Fichier de prospects B2B Paris (≥ 50 salariés), construit uniquement à partir
de sources publiques et sans deviner d'email nominatif.

Étapes :
  1. extract   : StockEtablissement + StockUniteLegale (SIRENE, data.gouv.fr)
                 → établissements actifs, CP 75001–75020 (+75116), tranche ≥ 21,
                   hors catégorie juridique 7xxx.
  2. site      : site officiel via Wikidata (propriété P1616 = SIREN, P856 = site).
  3. generiques: emails génériques (rh@, recrutement@, contact@…) publiés sur le
                 site de l'entreprise (robots.txt respecté).
  4. export    : fusion avec contacts_publics.csv (saisie manuelle, source
                 obligatoire) et écriture du tableur .xlsx.

Aucune donnée LinkedIn, aucun email nominatif reconstruit par motif.
"""
import argparse
import csv
import json
import re
import sys
import time
import urllib.robotparser
from pathlib import Path
from urllib.parse import urljoin, urlparse

SIRENE_ETAB_URL = "https://object.files.data.gouv.fr/data-pipeline-open/siren/stock/StockEtablissement_utf8.parquet"
SIRENE_UL_URL = "https://object.files.data.gouv.fr/data-pipeline-open/siren/stock/StockUniteLegale_utf8.parquet"
WIKIDATA_SPARQL = "https://query.wikidata.org/sparql"
USER_AGENT = "prospection-b2b-rgpd/1.0 (collecte de donnees publiques)"

CODES_POSTAUX = [f"750{i:02d}" for i in range(1, 21)]
CP_16E_BIS = "75116"  # second code postal du 16e arrondissement

TRANCHES = {
    "21": "50 à 99 salariés",
    "22": "100 à 199 salariés",
    "31": "200 à 249 salariés",
    "32": "250 à 499 salariés",
    "41": "500 à 999 salariés",
    "42": "1 000 à 1 999 salariés",
    "51": "2 000 à 4 999 salariés",
    "52": "5 000 à 9 999 salariés",
    "53": "10 000 salariés et plus",
}

# Préfixes d'emails génériques acceptés (jamais d'email nominatif).
GENERIQUES = (
    "rh", "drh", "hr", "recrutement", "recrutements", "recruitment", "recrute",
    "jobs", "job", "emploi", "emplois", "careers", "carrieres", "candidature",
    "candidatures", "talent", "talents", "contact", "info", "infos", "accueil",
    "bonjour", "hello", "office",
)
PAGES_CONTACT = (
    "", "contact", "contact/", "nous-contacter", "contactez-nous", "carrieres",
    "carrieres/", "careers", "recrutement", "rejoignez-nous", "mentions-legales",
    "equipe", "notre-equipe", "a-propos",
)
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")

COLONNES = [
    "Entreprise", "SIREN", "Effectif", "NAF", "Site", "Domaine", "Prénom", "Nom",
    "Poste", "Source", "Email public", "Email générique",
]
COLONNES_SUPPL = [
    "SIRET", "Adresse", "Code postal", "Siège", "Catégorie juridique",
    "Code tranche", "Année effectif", "Source site", "Source email générique",
    "Statut recherche",
]


# --------------------------------------------------------------------------- 1
def extract(etab_src, ul_src, out_csv, inclure_75116=True):
    import duckdb

    cps = CODES_POSTAUX + ([CP_16E_BIS] if inclure_75116 else [])
    con = duckdb.connect()
    if str(etab_src).startswith("http") or str(ul_src).startswith("http"):
        con.execute("INSTALL httpfs; LOAD httpfs;")

    def src(path):
        p = str(path)
        return f"read_parquet('{p}')" if p.endswith(".parquet") else f"read_csv('{p}', all_varchar=true)"

    cp_list = ",".join(f"'{c}'" for c in cps)
    tr_list = ",".join(f"'{t}'" for t in TRANCHES)
    query = f"""
    WITH e AS (
      SELECT lpad(CAST(siren AS VARCHAR), 9, '0') AS siren,
             lpad(CAST(siret AS VARCHAR), 14, '0') AS siret,
             lpad(CAST(codePostalEtablissement AS VARCHAR), 5, '0') AS cp,
             lpad(CAST(trancheEffectifsEtablissement AS VARCHAR), 2, '0') AS tranche,
             CAST(anneeEffectifsEtablissement AS VARCHAR) AS annee_eff,
             CAST(activitePrincipaleEtablissement AS VARCHAR) AS naf,
             CAST(etablissementSiege AS VARCHAR) AS siege,
             concat_ws(' ', coalesce(CAST(TRY_CAST(numeroVoieEtablissement AS BIGINT) AS VARCHAR),
                                CAST(numeroVoieEtablissement AS VARCHAR)),
                       indiceRepetitionEtablissement, typeVoieEtablissement,
                       libelleVoieEtablissement) AS voie,
             complementAdresseEtablissement AS complement,
             libelleCommuneEtablissement AS commune,
             coalesce(denominationUsuelleEtablissement, enseigne1Etablissement) AS enseigne
      FROM {src(etab_src)}
      WHERE etatAdministratifEtablissement = 'A'
        AND coalesce(statutDiffusionEtablissement, 'O') = 'O'
    ), u AS (
      SELECT lpad(CAST(siren AS VARCHAR), 9, '0') AS siren,
             coalesce(denominationUniteLegale,
                      trim(concat_ws(' ', prenom1UniteLegale, nomUniteLegale))) AS raison_sociale,
             CAST(categorieJuridiqueUniteLegale AS VARCHAR) AS cat_jur
      FROM {src(ul_src)}
      WHERE etatAdministratifUniteLegale = 'A'
    )
    SELECT u.raison_sociale, e.siren, e.siret, e.tranche, e.annee_eff, e.naf, e.siege,
           e.voie, e.complement, e.cp, e.commune, e.enseigne, u.cat_jur
    FROM e JOIN u USING (siren)
    WHERE e.cp IN ({cp_list})
      AND e.tranche IN ({tr_list})
      AND NOT starts_with(coalesce(u.cat_jur, ''), '7')
    ORDER BY e.cp, u.raison_sociale, e.siret
    """
    rows = con.execute(query).fetchall()
    cols = [d[0] for d in con.description]
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(cols)
        w.writerows(rows)
    print(f"[extract] {len(rows)} établissements → {out_csv}")
    return len(rows)


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# --------------------------------------------------------------------------- 2
def domaine_de(url):
    host = urlparse(url if "://" in url else "https://" + url).hostname or ""
    return host[4:] if host.startswith("www.") else host


def sites_wikidata(sirens, cache_path, batch=200):
    """SIREN → site officiel (P856) via Wikidata. Résultats mis en cache."""
    import requests

    cache = json.loads(Path(cache_path).read_text()) if Path(cache_path).exists() else {}
    todo = sorted({s for s in sirens if s not in cache})
    for i in range(0, len(todo), batch):
        chunk = todo[i:i + batch]
        values = " ".join(f'"{s}"' for s in chunk)
        q = f"""SELECT ?siren ?site ?item WHERE {{
                  VALUES ?siren {{ {values} }}
                  ?item wdt:P1616 ?siren . ?item wdt:P856 ?site . }}"""
        r = requests.get(WIKIDATA_SPARQL, params={"query": q, "format": "json"},
                         headers={"User-Agent": USER_AGENT}, timeout=60)
        r.raise_for_status()
        for s in chunk:
            cache.setdefault(s, None)
        for b in r.json()["results"]["bindings"]:
            s = b["siren"]["value"]
            if cache.get(s) is None:  # garde le premier site trouvé
                cache[s] = {"site": b["site"]["value"], "source": b["item"]["value"]}
        Path(cache_path).write_text(json.dumps(cache, ensure_ascii=False, indent=1))
        print(f"[site] {min(i + batch, len(todo))}/{len(todo)} SIREN interrogés")
        time.sleep(1)
    return cache


# --------------------------------------------------------------------------- 3
def est_generique(email, domaine):
    local, _, dom = email.lower().partition("@")
    if not (dom == domaine or dom.endswith("." + domaine)):
        return False
    return re.split(r"[._+-]", local)[0] in GENERIQUES


def emails_generiques(site, timeout=15, pause=1.0):
    """Emails génériques publiés sur le site (même domaine), avec l'URL source."""
    import requests

    base = site if site.endswith("/") else site + "/"
    domaine = domaine_de(site)
    rp = urllib.robotparser.RobotFileParser(urljoin(base, "/robots.txt"))
    try:
        rp.read()
    except Exception:
        rp = None
    trouves = {}
    for page in PAGES_CONTACT:
        url = urljoin(base, page)
        if rp is not None and not rp.can_fetch(USER_AGENT, url):
            continue
        try:
            r = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=timeout)
        except requests.RequestException:
            continue
        if r.status_code == 200 and "html" in r.headers.get("content-type", ""):
            texte = r.text.replace("&#64;", "@").replace("[at]", "@").replace("(at)", "@")
            for e in EMAIL_RE.findall(texte):
                e = e.lower().rstrip(".")
                if est_generique(e, domaine):
                    trouves.setdefault(e, url)
        time.sleep(pause)
    return trouves


def collecte_generiques(sites, cache_path):
    cache = json.loads(Path(cache_path).read_text()) if Path(cache_path).exists() else {}
    for n, (siren, site) in enumerate(sorted(sites.items()), 1):
        if siren in cache:
            continue
        cache[siren] = emails_generiques(site)
        Path(cache_path).write_text(json.dumps(cache, ensure_ascii=False, indent=1))
        print(f"[generiques] {n}/{len(sites)} {site} → {list(cache[siren]) or '-'}")
    return cache


# --------------------------------------------------------------------------- 4
def charge_contacts(path):
    """contacts_publics.csv : saisie manuelle. Une ligne sans source est rejetée."""
    if not path or not Path(path).exists():
        return {}
    contacts = {}
    for i, c in enumerate(read_csv(path), 2):
        if not c.get("siren"):
            continue
        if (c.get("nom") or c.get("email_public")) and not c.get("source"):
            print(f"[contacts] ligne {i} ignorée : source manquante", file=sys.stderr)
            continue
        contacts.setdefault(c["siren"].strip(), []).append(c)
    return contacts


def lignes_tableur(etabs, sites, generiques, contacts):
    lignes = []
    for e in etabs:
        siren = e["siren"]
        info_site = sites.get(siren) or {}
        site = info_site.get("site", "")
        gen = generiques.get(siren) or {}
        adresse = ", ".join(x for x in (e["voie"], e["complement"],
                                        f"{e['cp']} {e['commune']}") if x)
        nom = e["raison_sociale"] + (f" ({e['enseigne']})" if e.get("enseigne") else "")
        base = {
            "Entreprise": nom, "SIREN": siren,
            "Effectif": TRANCHES.get(e["tranche"], e["tranche"]), "NAF": e["naf"],
            "Site": site, "Domaine": domaine_de(site) if site else "",
            "Email générique": "; ".join(gen),
            "SIRET": e["siret"], "Adresse": adresse, "Code postal": e["cp"],
            "Siège": "oui" if str(e["siege"]).lower() in ("true", "1") else "non",
            "Catégorie juridique": e["cat_jur"], "Code tranche": e["tranche"],
            "Année effectif": e["annee_eff"], "Source site": info_site.get("source", ""),
            "Source email générique": "; ".join(sorted(set(gen.values()))),
        }
        personnes = contacts.get(siren) or [{}]
        for p in personnes:
            ligne = dict(base)
            ligne.update({
                "Prénom": p.get("prenom", ""), "Nom": p.get("nom", ""),
                "Poste": p.get("poste", ""), "Source": p.get("source", ""),
                "Email public": p.get("email_public", ""),
            })
            if p.get("email_generique"):
                ligne["Email générique"] = "; ".join(
                    x for x in (ligne["Email générique"], p["email_generique"]) if x)
            trouve = any(ligne[k] for k in ("Site", "Nom", "Email public", "Email générique"))
            ligne["Statut recherche"] = "Informations trouvées" if trouve else "Aucune information trouvée"
            lignes.append(ligne)
    return lignes


METHODO = [
    ("Source entreprises", "Base SIRENE (INSEE), fichiers StockEtablissement et StockUniteLegale, data.gouv.fr — Licence Ouverte 2.0."),
    ("Filtres", "Établissement actif (A) ; code postal 75001–75020 (+75116, 2e CP du 16e, sauf option --sans-75116) ; tranche d'effectif de l'ÉTABLISSEMENT ≥ 21 (50 salariés et +) ; catégorie juridique de l'unité légale hors 7xxx ; établissements en diffusion partielle exclus."),
    ("Effectif", "Tranche déclarée à l'INSEE pour l'année indiquée (colonne « Année effectif ») — peut être ancienne."),
    ("Site / Domaine", "Site officiel déclaré dans Wikidata (P856) pour le SIREN (P1616). Le domaine est déduit du site ; il n'est pas garanti identique au domaine email."),
    ("Email générique", "Adresses génériques (rh@, recrutement@, contact@…) publiées sur le site de l'entreprise, même domaine uniquement, robots.txt respecté. URL source en colonne « Source email générique »."),
    ("Contacts nominatifs", "Saisie manuelle depuis des sources publiques (site, page équipe, communiqués), URL source obligatoire. Aucune donnée LinkedIn. Aucun email nominatif deviné : colonne vide si non publié."),
    ("RGPD", "Base légale envisagée : intérêt légitime (art. 6.1.f). Informer chaque personne contactée dès le premier message (art. 14) : source des données, finalité, droit d'opposition. Prospection B2B par email : autorisée sans consentement si le message est en lien avec la fonction de la personne et propose une désinscription simple (CNIL). Durée de conservation : 3 ans max. après le dernier contact. Tenir une liste d'opposition."),
]


def export_xlsx(lignes, out_xlsx):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    entete = COLONNES + COLONNES_SUPPL

    def feuille(ws, rows):
        ws.append(entete)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for r in rows:
            ws.append([r.get(k, "") for k in entete])
        for i, k in enumerate(entete, 1):
            largeur = max([len(k)] + [len(str(r.get(k, ""))) for r in rows[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, largeur + 2), 60)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Prospects"
    feuille(ws, lignes)
    feuille(wb.create_sheet("Sans information"),
            [l for l in lignes if l["Statut recherche"] == "Aucune information trouvée"])
    m = wb.create_sheet("Méthodologie & RGPD")
    for row in METHODO:
        m.append(row)
    m.column_dimensions["A"].width = 22
    m.column_dimensions["B"].width = 140
    for c in m["A"]:
        c.font = Font(bold=True)
    wb.save(out_xlsx)
    sans = sum(l["Statut recherche"] == "Aucune information trouvée" for l in lignes)
    print(f"[export] {len(lignes)} lignes ({sans} sans information) → {out_xlsx}")


# --------------------------------------------------------------------------- CLI
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--etablissements", default=SIRENE_ETAB_URL, help="StockEtablissement (.parquet ou .csv, chemin ou URL)")
    ap.add_argument("--unites-legales", default=SIRENE_UL_URL, help="StockUniteLegale (.parquet ou .csv, chemin ou URL)")
    ap.add_argument("--sortie", default="out", help="dossier de travail et de sortie")
    ap.add_argument("--contacts", default="contacts_publics.csv", help="contacts saisis manuellement (sources publiques)")
    ap.add_argument("--sans-75116", action="store_true", help="exclure le CP 75116 (16e arrondissement)")
    ap.add_argument("--sans-site", action="store_true", help="ne pas interroger Wikidata")
    ap.add_argument("--sans-generiques", action="store_true", help="ne pas visiter les sites web")
    args = ap.parse_args(argv)

    out = Path(args.sortie)
    out.mkdir(parents=True, exist_ok=True)
    etab_csv = out / "etablissements_paris_50plus.csv"
    if not etab_csv.exists():
        extract(args.etablissements, args.unites_legales, etab_csv, not args.sans_75116)
    etabs = read_csv(etab_csv)
    sirens = {e["siren"] for e in etabs}

    sites = {} if args.sans_site else sites_wikidata(sirens, out / "cache_sites.json")
    sites = {s: v for s, v in sites.items() if v and s in sirens}
    generiques = {} if args.sans_generiques else collecte_generiques(
        {s: v["site"] for s, v in sites.items()}, out / "cache_generiques.json")

    lignes = lignes_tableur(etabs, sites, generiques, charge_contacts(args.contacts))
    export_xlsx(lignes, out / "prospects_paris_b2b.xlsx")


if __name__ == "__main__":
    main()
