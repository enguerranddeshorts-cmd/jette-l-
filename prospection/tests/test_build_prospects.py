"""Test du pipeline sur des données SYNTHÉTIQUES (fictives) — aucune donnée réelle."""
import csv
import sys
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import build_prospects as bp  # noqa: E402


def _etab(siren, nic, cp, tranche, etat="A", diff="O"):
    return dict(siren=siren, siret=siren + nic, etatAdministratifEtablissement=etat,
                statutDiffusionEtablissement=diff, codePostalEtablissement=cp,
                trancheEffectifsEtablissement=tranche, anneeEffectifsEtablissement="2023",
                activitePrincipaleEtablissement="62.01Z", etablissementSiege=True,
                numeroVoieEtablissement="12", indiceRepetitionEtablissement=None,
                typeVoieEtablissement="RUE", libelleVoieEtablissement="DE TEST",
                complementAdresseEtablissement=None, libelleCommuneEtablissement="PARIS",
                enseigne1Etablissement=None, denominationUsuelleEtablissement=None)


def _ul(siren, nom, cj):
    return dict(siren=siren, denominationUniteLegale=nom, prenom1UniteLegale=None,
                nomUniteLegale=None, categorieJuridiqueUniteLegale=cj,
                etatAdministratifUniteLegale="A")


def test_pipeline(tmp_path):
    etabs = [
        _etab("100000001", "00011", "75008", "21"),          # gardé
        _etab("100000002", "00012", "75116", "53"),          # gardé (75116)
        _etab("100000003", "00013", "75009", "12"),          # < 50 salariés
        _etab("100000004", "00014", "92100", "41"),          # hors Paris
        _etab("100000005", "00015", "75001", "22"),          # administration (7xxx)
        _etab("100000006", "00016", "75002", "32", etat="F"),  # fermé
        _etab("100000007", "00017", "75003", "NN"),          # tranche inconnue
        _etab("100000008", "00018", "75004", "31", diff="P"),  # diffusion partielle
    ]
    uls = [_ul("10000000" + str(i), f"SOCIETE FICTIVE {i}", "7210" if i == 5 else "5710")
           for i in range(1, 9)]
    pq.write_table(pa.Table.from_pylist(etabs), tmp_path / "e.parquet")
    pq.write_table(pa.Table.from_pylist(uls), tmp_path / "u.parquet")
    (tmp_path / "contacts.csv").write_text(
        "siren,prenom,nom,poste,source,email_public,email_generique,date_collecte\n"
        "100000001,Alex,Exemple,DRH,https://exemple.invalid/equipe,,,2026-10-14\n"
        "100000002,Sans,Source,RRH,,,,2026-10-14\n")

    bp.main(["--etablissements", str(tmp_path / "e.parquet"),
             "--unites-legales", str(tmp_path / "u.parquet"),
             "--sortie", str(tmp_path / "out"), "--contacts", str(tmp_path / "contacts.csv"),
             "--sans-site", "--sans-generiques"])

    rows = list(csv.DictReader(open(tmp_path / "out/etablissements_paris_50plus.csv")))
    assert sorted(r["siren"] for r in rows) == ["100000001", "100000002"]
    assert rows[0]["voie"] == "12 RUE DE TEST"

    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "out/prospects_paris_b2b.xlsx")
    ws = wb["Prospects"]
    head = [c.value for c in ws[1]]
    assert head[:12] == bp.COLONNES
    data = [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    a = next(d for d in data if d["SIREN"] == "100000001")
    assert a["Nom"] == "Exemple" and a["Source"] and a["Email public"] is None
    b = next(d for d in data if d["SIREN"] == "100000002")
    assert b["Nom"] is None  # contact sans source rejeté
    assert b["Statut recherche"] == "Aucune information trouvée"
    assert wb["Sans information"].max_row == 2


def test_est_generique():
    assert bp.est_generique("recrutement@acme.fr", "acme.fr")
    assert bp.est_generique("rh.paris@groupe.acme.fr", "acme.fr")
    assert not bp.est_generique("jean.dupont@acme.fr", "acme.fr")
    assert not bp.est_generique("contact@autre.fr", "acme.fr")
    assert bp.domaine_de("https://www.acme.fr/fr/") == "acme.fr"
